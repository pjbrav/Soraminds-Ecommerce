"""Template generator (spec checklist #2).

Regenerates the owner-facing .xlsx upload template for a tenant, with
data-validation dropdowns sourced from the tenant's Category Reference
(config YAML), mirroring the shipped SpiceHub template's structure:
  Sheet 1: Instructions
  Sheet 2: Category Reference
  Sheet 3: Menu Items (headers + hint row)

Usage:  python scripts/generate_template.py [tenant_id] [output.xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.tenant_config import get_tenant

HEADERS = [
    "item_id", "category", "subcategory", "item_name", "description", "price",
    "dietary_type", "spice_level", "calorie_range", "allergens", "image_filename",
    "is_featured", "is_available", "sort_order", "tags",
]
HINTS = [
    "Auto-generated if left blank", "Required — pick from dropdown",
    "Optional — free text", "Required", "Optional — auto-generated if blank",
    "Required — number only", "Required — pick from dropdown",
    "Optional — defaults to None", "Optional — e.g. 350-500",
    "Optional — comma-separated", "Required — must match uploaded photo",
    "Optional — defaults to No", "Optional — defaults to Yes",
    "Optional — auto-assigned if blank", "Optional — comma-separated",
]
INSTRUCTIONS = [
    ("SpiceHub Kitchen — Menu Upload Template", None),
    ("How to fill this file out — read this sheet first, then go to the “Menu Items” tab.", None),
    (None, None),
    ("Which sheet do I edit?", "Only the “Menu Items” sheet. Do not edit “Instructions” or “Category Reference.”"),
    ("One row = one menu item", "Add a new row for every dish, drink, or side you sell. Do not merge cells or leave blank rows between items."),
    ("Required fields", "category, item_name, price, dietary_type, and image_filename must be filled in for every row. The upload will be rejected otherwise."),
    ("Dropdowns", "category, dietary_type, spice_level, is_featured, and is_available are dropdown menus. Click the cell and choose from the list — do not type a value that isn’t on the list."),
    ("Images", "Upload your photos separately as a ZIP file. In image_filename, type the exact file name of the matching photo (e.g., goat_biryani.jpg). File names are not case-sensitive, but must otherwise match exactly."),
    ("Descriptions", "If you leave description blank, one will be generated for you automatically. You can review and edit it after upload before publishing."),
    ("Prices", "Numbers only, no currency symbol (e.g., 13.99, not $13.99)."),
    ("is_featured", "Set to Yes for the handful of items you want highlighted at the top of your kiosk menu — most restaurants feature 4-8 items."),
    ("What you don’t need to fill in", "“Most Liked” badges and like-percentages are calculated automatically from real order history once you’re live — there’s no column for them."),
    ("Optional fields", "subcategory, calorie_range, allergens, sort_order, and tags can be left blank — the system will apply sensible defaults."),
    ("What happens after upload?", "You’ll see a validation report showing any rows that need a fix, with a suggested correction where possible. Nothing goes live until you review the preview and click Publish."),
]


def generate(tenant_id: str, output: Path) -> Path:
    tenant = get_tenant(tenant_id)
    wb = Workbook()

    instructions = wb.active
    instructions.title = "Instructions"
    for row in INSTRUCTIONS:
        instructions.append(list(row))

    ref = wb.create_sheet("Category Reference")
    ref.append(["Valid Categories"])
    ref.append(["Matches the restaurant's configured menu structure. Pick the closest match."])
    ref.append(["Category", "Examples", None, "Dietary Type", "Spice Level", "Yes / No"])
    examples = {
        "Veg Appetizers": "Samosa, Gobi Manchurian, Chilli Paneer",
        "Non Veg Appetizers": "Chicken 65, Tandoori Shrimp, Chicken Lollipop",
        "Veg Curries": "Paneer Butter Masala, Dal Makhani, Malai Kofta",
        "Non Veg Curries": "Butter Chicken, Egg Masala, Chettinad Goat Curry",
        "Rice": "Veg Fried Rice, Chicken Fried Rice",
        "Hyderabadi Dum Biryani": "Chicken Biryani, Goat Biryani, Paneer Biryani",
        "Breads": "Naan, Garlic Naan, Roti",
        "Desserts": "Gulab Jamun, Rasmalai",
        "Drinks": "Mango Lassi",
        "Soups": "Veg Soup, Sweet Corn Soup",
    }
    for i, category in enumerate(tenant.categories):
        col = [category, examples.get(category, "")]
        if i == 0:
            col += [None, *tenant.dietary_types[:1], *tenant.spice_levels[:1], "Yes"]
        elif i == 1:
            col += [None, *tenant.dietary_types[1:2], *tenant.spice_levels[1:2], "No"]
        elif i == 2:
            col += [None, *tenant.dietary_types[2:3], *tenant.spice_levels[2:3]]
        elif i == 3:
            col += [None, *tenant.dietary_types[3:4], *tenant.spice_levels[3:4]]
        elif i == 4:
            col += [None, None, *tenant.spice_levels[4:5]]
        ref.append(col)

    menu = wb.create_sheet("Menu Items")
    menu.append(HEADERS)
    menu.append(HINTS)
    for _ in range(30):  # blank rows for the owner to fill in
        menu.append([None] * len(HEADERS))

    last_row = menu.max_row
    cat_dv = DataValidation(
        type="list",
        formula1=f"='Category Reference'!$A$4:$A${3 + len(tenant.categories)}",
        allow_blank=True,
    )
    diet_dv = DataValidation(type="list", formula1='"' + ",".join(tenant.dietary_types) + '"', allow_blank=True)
    spice_dv = DataValidation(type="list", formula1='"' + ",".join(tenant.spice_levels) + '"', allow_blank=True)
    yesno_dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    for dv, column in ((cat_dv, "B"), (diet_dv, "G"), (spice_dv, "H"), (yesno_dv, "L"), (yesno_dv, "M")):
        menu.add_data_validation(dv)
        dv.add(f"{column}3:{column}{last_row}")

    for i, header in enumerate(HEADERS, start=1):
        menu.cell(row=1, column=i).font = menu.cell(row=1, column=i).font.copy(bold=True)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


if __name__ == "__main__":
    tenant_id = sys.argv[1] if len(sys.argv) > 1 else "spicehub-kitchen-troy"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("generated_template.xlsx")
    print(f"Wrote {generate(tenant_id, out)}")
